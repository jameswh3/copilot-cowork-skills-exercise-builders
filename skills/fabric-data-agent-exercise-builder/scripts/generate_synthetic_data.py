#!/usr/bin/env python3
"""Deterministic synthetic data + expected-answer generator for Fabric data agent exercises.

Reads a scenario spec (JSON), generates related tables with consistent keys and
subtle, realistic data-quality traps, writes one CSV per table (a genuine flat
export for OneLake upload), optionally one .xlsx workbook (one tab per table,
Aptos font), and computes every expected answer by executing SQL over the
generated data in an in-memory SQLite database. Nothing is estimated.

Outputs:
  --out  <table>.csv               one per table, UTF-8, header row, ISO dates (participant files)
  --out  <workbook>.xlsx           optional (--xlsx NAME), one tab per table (participant file)
  --meta manifest.json             schemas, row counts, lakehouse folder; input to build_ingest_notebook.py
  --meta expected_answers.json     every answer and example query, executed, with result rows
  (--meta files are authoring-only and never ship to participants)

Usage:
  python generate_synthetic_data.py --spec spec.json --out working/<slug>/data [--xlsx <Name>.xlsx]
  python generate_synthetic_data.py --spec spec.json --check     # validate only, write nothing

Generator kinds (column "gen"):
  sequence  {prefix, start, width}          id strings such as WO-004512
  list      {values}                         cycles values in order (row i -> values[i % n])
  choice    {values, weights?}               weighted random pick
  int       {min, max}
  float     {min, max, decimals?}
  date      {start, end}                     ISO yyyy-mm-dd
  timestamp {start, end}                     ISO yyyy-mm-dd HH:MM:SS
  date_offset {from, min_days, max_days, tail?: {rate, min_days, max_days}}
  fk        {table, column}                  random existing key from a parent table
  lookup    {via, table, key, column}        copies a parent attribute through an fk column
  const     {value}

Column options (applied in this order after generation):
  nullify_if  {column, in: [...]}            set NULL when another column's canonical value is listed
  (table) supersede {rate, key, version_column, current_flag, mutate: {col: gen}}
  variants    {canonical: [variant, ...], rate}   inconsistent codes for the same meaning
  unit_trap   {unit_column, alt_unit, factor, rate}  value stored in an alternate unit
  null_rate   float                          random missing values

Only the Python standard library is required; openpyxl is needed only for --xlsx.
"""
import argparse
import csv
import datetime as dt
import json
import os
import random
import re
import sqlite3
import sys

TYPES = {"string", "int", "double", "date", "timestamp", "boolean"}


def _d(s):
    return dt.date.fromisoformat(s)


def _ts(s):
    return dt.datetime.fromisoformat(s)


class Generator:
    def __init__(self, spec):
        self.spec = spec
        self.rng = random.Random(spec.get("seed", 20261022))
        self.data = {}      # table -> list of dict rows
        self.schemas = {}   # table -> list of (name, type)

    # ---------- value generation ----------
    def gen_value(self, g, i, row):
        k = g["kind"]
        r = self.rng
        if k == "sequence":
            return f"{g.get('prefix', '')}{g.get('start', 1) + i:0{g.get('width', 1)}d}"
        if k == "list":
            vals = g["values"]
            return vals[i % len(vals)]
        if k == "choice":
            return r.choices(g["values"], weights=g.get("weights"))[0]
        if k == "int":
            return r.randint(g["min"], g["max"])
        if k == "float":
            return round(r.uniform(g["min"], g["max"]), g.get("decimals", 2))
        if k == "date":
            a, b = _d(g["start"]), _d(g["end"])
            return a + dt.timedelta(days=r.randint(0, (b - a).days))
        if k == "timestamp":
            a, b = _ts(g["start"]), _ts(g["end"])
            secs = int((b - a).total_seconds())
            return (a + dt.timedelta(seconds=r.randint(0, secs))).replace(microsecond=0)
        if k == "date_offset":
            base = row.get(g["from"])
            if base is None:
                return None
            tail = g.get("tail")
            if tail and r.random() < tail["rate"]:
                days = r.randint(tail["min_days"], tail["max_days"])
            else:
                days = r.randint(g["min_days"], g["max_days"])
            return base + dt.timedelta(days=days)
        if k == "fk":
            parent = self.data[g["table"]]
            return r.choice(parent)[g["column"]]
        if k == "lookup":
            key_val = row[g["via"]]
            for p in self.data[g["table"]]:
                if p[g["key"]] == key_val:
                    return p[g["column"]]
            return None
        if k == "const":
            return g["value"]
        raise ValueError(f"unknown generator kind: {k}")

    # ---------- table build ----------
    def build_table(self, t):
        name = t["name"]
        cols = t["columns"]
        for c in cols:
            if c["type"] not in TYPES:
                raise ValueError(f"{name}.{c['name']}: type must be one of {sorted(TYPES)}")
        self.schemas[name] = [(c["name"], c["type"]) for c in cols]
        rows = []
        for i in range(t["rows"]):
            row = {}
            for c in cols:
                row[c["name"]] = self.gen_value(c["gen"], i, row)
            rows.append(row)
        # nullify_if (on canonical values)
        for c in cols:
            rule = c.get("nullify_if")
            if rule:
                for row in rows:
                    if row.get(rule["column"]) in rule["in"]:
                        row[c["name"]] = None
        # supersede: append later revisions of a share of records
        sup = t.get("supersede")
        if sup:
            extra = []
            for row in rows:
                if self.rng.random() < sup["rate"]:
                    new = dict(row)
                    new[sup["version_column"]] = row[sup["version_column"]] + 1
                    row[sup["current_flag"]] = False
                    new[sup["current_flag"]] = True
                    for col, g in sup.get("mutate", {}).items():
                        new[col] = self.gen_value(g, 0, new)
                    extra.append(new)
            rows.extend(extra)
            rows.sort(key=lambda x: (str(x[sup["key"]]), x[sup["version_column"]]))
        # variants, unit traps, nulls
        for c in cols:
            var = c.get("variants")
            if var:
                for row in rows:
                    v = row[c["name"]]
                    if v in var["canonical"] and self.rng.random() < var["rate"]:
                        row[c["name"]] = self.rng.choice(var["canonical"][v])
            ut = c.get("unit_trap")
            if ut:
                for row in rows:
                    if row[c["name"]] is not None and self.rng.random() < ut["rate"]:
                        row[c["name"]] = round(row[c["name"]] * ut["factor"], 2)
                        row[ut["unit_column"]] = ut["alt_unit"]
            nr = c.get("null_rate")
            if nr:
                for row in rows:
                    if self.rng.random() < nr:
                        row[c["name"]] = None
        self.data[name] = rows

    def build(self):
        for t in self.spec["tables"]:
            self.build_table(t)

    # ---------- checks ----------
    def key_integrity(self):
        problems = []
        for t in self.spec["tables"]:
            for c in t["columns"]:
                g = c["gen"]
                if g["kind"] == "fk":
                    parent_keys = {p[g["column"]] for p in self.data[g["table"]]}
                    bad = [r[c["name"]] for r in self.data[t["name"]]
                           if r[c["name"]] is not None and r[c["name"]] not in parent_keys]
                    if bad:
                        problems.append(f"{t['name']}.{c['name']}: {len(bad)} orphan keys")
        return problems


def fmt(v, typ):
    if v is None:
        return ""
    if typ == "boolean":
        return "true" if v else "false"
    if typ == "timestamp":
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if typ == "date":
        return v.isoformat()
    return v


def sqlite_val(v, typ):
    if v is None:
        return None
    if typ == "boolean":
        return 1 if v else 0
    if typ in ("date", "timestamp"):
        return fmt(v, typ)
    return v


SQLITE_TYPE = {"string": "TEXT", "int": "INTEGER", "double": "REAL", "date": "TEXT",
               "timestamp": "TEXT", "boolean": "INTEGER"}


def load_sqlite(gen):
    con = sqlite3.connect(":memory:")
    for name, schema in gen.schemas.items():
        cols = ", ".join(f'"{c}" {SQLITE_TYPE[t]}' for c, t in schema)
        con.execute(f'CREATE TABLE "{name}" ({cols})')
        ph = ", ".join("?" for _ in schema)
        con.executemany(f'INSERT INTO "{name}" VALUES ({ph})',
                        [[sqlite_val(r[c], t) for c, t in schema] for r in gen.data[name]])
    return con


FROM_RE = re.compile(r"\b(?:FROM|JOIN)\s+(?:\[?dbo\]?\.)?\[?([A-Za-z_][A-Za-z0-9_]*)\]?", re.I)
QUAL_RE = re.compile(r"\b([A-Za-z_][A-Za-z0-9_]*)\.\[?([A-Za-z_][A-Za-z0-9_]*)\]?")


def static_check_tsql(sql, schemas):
    """Light schema check of participant-facing T-SQL: tables after FROM/JOIN exist and
    alias.column references resolve to a real column. Not a parser; catches typos."""
    issues = []
    tables = set(schemas)
    all_cols = {c for s in schemas.values() for c, _ in s}
    for t in FROM_RE.findall(sql):
        if t not in tables:
            issues.append(f"unknown table '{t}'")
    for alias, col in QUAL_RE.findall(sql):
        if alias.lower() == "dbo" or alias in tables and col in tables:
            continue
        if col not in all_cols and col not in tables:
            issues.append(f"unknown column '{alias}.{col}'")
    return issues


def run_sql(con, sql):
    cur = con.execute(sql)
    cols = [d[0] for d in cur.description]
    rows = [list(r) for r in cur.fetchall()]
    return cols, rows


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--spec", required=True)
    ap.add_argument("--out")
    ap.add_argument("--xlsx", help="optional workbook filename, one tab per table")
    ap.add_argument("--meta", help="folder for manifest.json and expected_answers.json "
                    "(authoring-only; default: <out>/../authoring, never shipped to participants)")
    ap.add_argument("--check", action="store_true", help="validate spec and SQL only")
    a = ap.parse_args()

    with open(a.spec, encoding="utf-8") as f:
        spec = json.load(f)
    gen = Generator(spec)
    gen.build()

    problems = gen.key_integrity()
    con = load_sqlite(gen)

    answers = []
    for item in spec.get("answers", []):
        try:
            cols, rows = run_sql(con, item["sql"])
            answers.append({**item, "columns": cols, "rows": rows})
        except sqlite3.Error as e:
            problems.append(f"answer {item.get('id')}: {e}")

    examples = []
    for ex in spec.get("example_queries", []):
        issues = static_check_tsql(ex["tsql"], gen.schemas)
        rec = {**ex, "static_issues": issues}
        check_sql = ex.get("sqlite", ex["tsql"] if ex.get("portable", True) else None)
        if check_sql:
            try:
                cols, rows = run_sql(con, check_sql)
                rec.update(columns=cols, rows=rows[:25])
            except sqlite3.Error as e:
                problems.append(f"example '{ex['question'][:40]}': {e}")
        problems += [f"example '{ex['question'][:40]}': {i}" for i in issues]
        examples.append(rec)

    counts = {n: len(r) for n, r in gen.data.items()}
    summary = {"ok": not problems, "problems": problems, "row_counts": counts}

    if a.check or not a.out:
        print(json.dumps(summary, indent=2))
        sys.exit(0 if not problems else 1)

    os.makedirs(a.out, exist_ok=True)
    files = []
    for name, schema in gen.schemas.items():
        path = os.path.join(a.out, f"{name}.csv")
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow([c for c, _ in schema])
            for r in gen.data[name]:
                w.writerow([fmt(r[c], t) for c, t in schema])
        files.append(f"{name}.csv")

    if a.xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        wb.remove(wb.active)
        body, head = Font(name="Aptos", size=11), Font(name="Aptos", size=11, bold=True)
        for name, schema in gen.schemas.items():
            ws = wb.create_sheet(title=name[:31])
            ws.append([c for c, _ in schema])
            for r in gen.data[name]:
                ws.append([fmt(r[c], t) if t in ("date", "timestamp") else r[c] for c, t in schema])
            for row in ws.iter_rows():
                for cell in row:
                    cell.font = head if cell.row == 1 else body
        wb.save(os.path.join(a.out, a.xlsx))
        files.append(a.xlsx)

    manifest = {
        "company": spec.get("company"),
        "scenario_slug": spec.get("scenario_slug"),
        "lakehouse_folder": spec.get("lakehouse_folder", "Files/raw"),
        "seed": spec.get("seed", 20261022),
        "files": files,
        "tables": [{"name": n, "csv": f"{n}.csv", "row_count": counts[n],
                    "columns": [{"name": c, "type": t} for c, t in s]}
                   for n, s in gen.schemas.items()],
    }
    meta = a.meta or os.path.join(os.path.dirname(os.path.abspath(a.out)), "authoring")
    os.makedirs(meta, exist_ok=True)
    with open(os.path.join(meta, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    with open(os.path.join(meta, "expected_answers.json"), "w", encoding="utf-8") as f:
        json.dump({"answers": answers, "example_queries": examples}, f, indent=2, default=str)
    print(json.dumps(summary, indent=2))
    sys.exit(0 if not problems else 1)


if __name__ == "__main__":
    main()
